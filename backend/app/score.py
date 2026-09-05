"""Compare a run against the human's own answers.

A development benchmark, and **deliberately not a tool the agent can reach.**
An agent that can see the answer key is marking its own homework, and every
number it then produces is worthless as evidence. Nothing in `agent.py`,
`tools.py` or either kit imports this module, and nothing should.

What it measures, and what it does not:

- **Resolution** has an answer key, so agreement is a real number. The workbook
  records what a person pulled out of each narrative and what they matched it
  to, and those columns are directly comparable.
- **Classification is judgement**, and disagreement is not automatically an
  error — the workbook itself carries rows a person marked `Review`. It is
  reported as agreement, never as accuracy.

Rows are joined on `Balance`, which is unique across all 100 rows in the
supplied week. Joining on position would silently mis-align the moment a parse
dropped or added a row, and produce a confident score for a comparison that
never happened.
"""

from __future__ import annotations

import json
from decimal import Decimal, InvalidOperation
from pathlib import Path

from app.kit.reference_kit import normalise
from app.reference.tables import resolve_source

STAGING = "Staging Sheet"


def _decimal(value: object) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        return Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return None


def load_truth(location: str) -> dict[Decimal, dict]:
    """The human's answers, keyed on the balance that identifies each row."""
    import openpyxl

    book = openpyxl.load_workbook(resolve_source(location), read_only=True, data_only=True)
    try:
        sheet = book[STAGING]
        rows = list(sheet.iter_rows(values_only=True))
        headers = [normalise(c) for c in rows[0]]
        index = {name: i for i, name in enumerate(headers) if name}

        truth = {}
        for cells in rows[1:]:
            if not any(c is not None for c in cells):
                continue
            balance = _decimal(cells[index["Balance"]])
            if balance is None:
                continue
            truth[balance] = {
                "counterparty_raw": normalise(cells[index["Pulled Out Sender/Beneficiary"]]),
                "counterparty_matched": normalise(cells[index["Matched Sender/Beneficiary"]]),
                "project_raw": normalise(cells[index["Pulled Out Project Code"]]),
                "project_matched": normalise(cells[index["Matched Project Code"]]),
                "classification": normalise(cells[index["Classification"]]),
            }
        return truth
    finally:
        book.close()


def _resolution(row: dict, field: str) -> dict:
    value = row.get(field)
    return value if isinstance(value, dict) else {}


def score_rows(rows: list[dict], truth: dict[Decimal, dict]) -> dict:
    """Agreement between one run's rows and the human's, field by field."""
    tally = {
        "rows": len(rows),
        "joined": 0,
        "unjoined": 0,
        "counterparty": {"both_named": 0, "agent_only": 0, "human_only": 0, "neither": 0},
        "counterparty_matched": {"agree": 0, "differ": 0, "agent_only": 0, "human_only": 0},
        "project_matched": {"agree": 0, "differ": 0, "agent_only": 0, "human_only": 0},
        "classification": {"agree": 0, "differ": 0},
        "disagreements": [],
    }

    for row in rows:
        balance = _decimal(row.get("balance"))
        want = truth.get(balance) if balance is not None else None
        if want is None:
            tally["unjoined"] += 1
            continue
        tally["joined"] += 1

        # Did each side pull a name out of the narrative at all?
        mine = bool(normalise(row.get("counterparty_raw")))
        theirs = bool(want["counterparty_raw"])
        key = (
            "both_named" if mine and theirs
            else "agent_only" if mine
            else "human_only" if theirs
            else "neither"
        )
        tally["counterparty"][key] += 1

        for field, column in (
            ("counterparty_match", "counterparty_matched"),
            ("project_code_match", "project_matched"),
        ):
            got = normalise(_resolution(row, field).get("matched_name"))
            expected = want[column]
            # The workbook writes an explicit sentinel where a project code did
            # not resolve. That is a "no match", not a value to agree with.
            if expected.lower().startswith("flag for review"):
                expected = ""
            if got and expected:
                same = got.casefold() == expected.casefold()
                tally[column]["agree" if same else "differ"] += 1
                if not same:
                    tally["disagreements"].append(
                        {"balance": str(balance), "field": column, "agent": got, "human": expected}
                    )
            elif got:
                tally[column]["agent_only"] += 1
            elif expected:
                tally[column]["human_only"] += 1

        got = normalise(row.get("classification"))
        expected = want["classification"]
        same = got.casefold() == expected.casefold()
        tally["classification"]["agree" if same else "differ"] += 1
        if not same and got and expected:
            tally["disagreements"].append(
                {"balance": str(balance), "field": "classification", "agent": got, "human": expected}
            )

    return tally


def score_runs(paths: list[Path], location: str, holdout: list[str] | None = None) -> dict:
    """Score several runs against one answer key, split tune from hold-out.

    The split is the whole point. Prompts are changed by looking at failures,
    which means a number measured on the documents that were looked at is a
    number about this week's data rather than about the pipeline. So the
    documents named in `holdout` are never opened while tuning, and **their
    number is the result**; the tune number is reported beside it only so the
    gap between them is visible. A wide gap means we memorised.
    """
    truth = load_truth(location)
    holdout = holdout or []
    per_run = []
    totals: dict[str, dict | None] = {"tune": None, "holdout": None, "all": None}

    for path in paths:
        payload = json.loads((path / "rows.json").read_text(encoding="utf-8"))
        account = payload.get("account", "")
        result = score_rows(payload.get("rows", []), truth)
        result["run"] = path.name
        result["account"] = account
        result["split"] = "holdout" if account in holdout else "tune"
        per_run.append(result)

        for key in (result["split"], "all"):
            totals[key] = result if totals[key] is None else _add(totals[key], result)

    return {
        "truth_rows": len(truth),
        "holdout_accounts": holdout,
        "runs": per_run,
        "tune": totals["tune"] or {},
        "holdout": totals["holdout"] or {},
        "total": totals["all"] or {},
    }


def _add(a: dict, b: dict) -> dict:
    out = {}
    for key, value in a.items():
        other = b.get(key)
        if isinstance(value, int):
            out[key] = value + (other or 0)
        elif isinstance(value, dict):
            out[key] = {k: v + other.get(k, 0) for k, v in value.items()}
        elif isinstance(value, list):
            out[key] = value + (other or [])
        else:
            out[key] = value
    return out
